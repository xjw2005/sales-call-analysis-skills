"""把飞书多维表导出的 Excel 导入数据库。

    python scripts/import_feishu_xlsx.py --file 导出.xlsx            # 只试运行，出报告，不写库
    python scripts/import_feishu_xlsx.py --file 导出.xlsx --commit   # 确认报告没问题后，真正写库

- 可以重复运行：门店按「门店编号 ST-xxxx」、拜访按「拜访编号 RA-xxxx」、总门店清单按「平台 + 门店 ID」更新，不会重复插入。
- 导入内容：人员 → 总门店清单 → 门店主档（含档案七维度）→ 拜访（录音文件名、转写、分析原文、DSR 纠正）→ 后续行动（月度目标）。
- 旧的分析结果是排好版的文字，按原文保存，页面按文本展示（拜访的 legacy 标记为真）。
- Excel 里只有录音文件名，没有音频本身：录音段落标记为 file_missing，音频需另行从飞书下载后放到服务器。
- 里面有真实电话和对话内容：不要把 Excel 提交到 Git，在服务器上运行。
"""

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy import select  # noqa: E402
from sqlalchemy.orm import Session  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.models import (  # noqa: E402
    Store, StoreCorrection, StoreDirectory, StoreProfileSection, Todo, User, Visit, VisitSegment,
)
from app.services.constants import (  # noqa: E402
    COOPERATION_STATUSES, FIRST_STAGE, LEGACY_STAGE_MAP, PROFILE_SECTIONS, STAGES, STORE_CONDITIONS,
)
from app.services.ingest import upsert_module, upsert_transcript  # noqa: E402

SHEET_DIRECTORY = "智生活+西港总门店清单（含企微）"
SHEET_STORES = "01 门店主档"
SHEET_VISITS = "02 门店拜访记录"
SHEET_TODOS = "03 后续行动"
ALL_PARTS = ["directory", "stores", "visits", "todos"]

# 主档里档案七维度的列名 → 档案维度 key
PROFILE_COLUMNS = {
    "一句画像": "one_line", "门店基本信息": "basic", "主营品类与品牌": "categories_brands", "经营模式": "business_model",
    "选品偏好": "selection_motion", "利润偏好": "price_profit", "合作偏好与排斥项": "cooperation_preferences",
}
PROFILE_LABEL_BY_KEY = dict(PROFILE_SECTIONS)
# 02 表里分析结果的列 → (模块, 证据所在的列)
ANALYSIS_COLUMNS = [
    ("AI拜访摘要", "ai-summary", None),
    ("显性需求(仅供参考)", "explicit-needs", "显性需求_原句参考"),
    ("隐性需求(仅供参考)", "implicit-needs", "隐性需求_原句参考"),
    ("关心类目打标", "concerns", "场景化类目归因"),
    ("合作进展打分评估", "effectiveness", "原句参考_合作进展打分评估"),
    ("核心金句", "quotes", None),
    ("门店档案", "store-profile", "门店档案-原文证据"),
    ("下一步行动策略", "next-action", None),
    ("上一次行动与这一次行动总结闭环", "loop", None),
]
SURVEY_COLUMNS = {
    "店铺面积": "area", "店铺全品类月均销售（罐）": "monthlyTotal", "店铺类型": "storeType", "是否做跨境": "crossBorder",
    "达能月均预估销量（罐）": "danoneMonthlyEst", "a2月均预估销售（罐）": "a2MonthlyEst",
}
REQUIRED = {
    SHEET_DIRECTORY: ["门店名称", "销售平台", "门店ID"],
    SHEET_STORES: ["门店编号", "门店名称", "装机平台", "门店ID", "合作状态", "归属销售"],
    SHEET_VISITS: ["拜访编号", "门店编号", "拜访区域经理", "进店时间", "拜访阶段1"],
    SHEET_TODOS: ["行动事项", "门店名称"],
}

TS = r"(?:\d+:)?\d+:\d+(?:\.\d+)?"
LINE_RE = re.compile(rf"^\[({TS})\s*[–—~\-]\s*({TS})\]\s*(?:(销售|客户|旁人)\s*[:：]\s*)?(.*)$")


# ---------- 小工具 ----------
def s(x) -> str:
    """单元格 → 去掉首尾空白的字符串；空返回 ''"""
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return str(x).strip()


def num(x) -> float | None:
    try:
        return float(s(x)) if s(x) else None
    except ValueError:
        return None


def yes(x) -> bool | None:
    return True if s(x) == "是" else (False if s(x) == "否" else None)


def ts_ms(t: str) -> int:
    parts = t.split(":")
    secs = 0.0
    for p in parts:
        secs = secs * 60 + float(p)
    return int(round(secs * 1000))


def parse_transcript(text: str) -> list[dict]:
    """`[00:11.450–00:12.290] 客户：……` → 逐句结构；没有时间戳的行接在上一句后面"""
    out: list[dict] = []
    for line in text.split("\n"):
        if not line.strip():
            continue
        m = LINE_RE.match(line.strip())
        if m:
            out.append({"uid": f"U{len(out) + 1:04d}", "role": m.group(3) or "", "startMs": ts_ms(m.group(1)),
                        "endMs": ts_ms(m.group(2)), "text": m.group(4).strip()})
        elif out:
            out[-1]["text"] += "\n" + line.strip()
    return out


def read_sheet(wb, name: str, report: dict) -> list[dict]:
    if name not in wb.sheetnames:
        report["warnings"].append(f"Excel 里没有名为「{name}」的表，已跳过")
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    header = [s(h) for h in rows[0]]
    missing = [c for c in REQUIRED[name] if c not in header]
    if missing:
        raise SystemExit(f"「{name}」缺少必需的列：{'、'.join(missing)}（导出时列名被改过？）")
    return [dict(zip(header, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


def to_utc(v, tz: ZoneInfo | None) -> datetime | None:
    """Excel 里的时间是不带时区的：按 --tz 解释后统一转成 UTC 存库"""
    if not isinstance(v, datetime):
        return None
    if tz is None:
        return v.replace(tzinfo=None)
    return v.replace(tzinfo=tz).astimezone(timezone.utc).replace(tzinfo=None)


def parse_date(x) -> date | None:
    if isinstance(x, datetime):
        return x.date()
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s(x))
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


# ---------- 各部分导入 ----------
def get_users(db: Session, names: set[str], role: str, report: dict) -> dict[str, User]:
    users = {u.name: u for u in db.scalars(select(User))}
    rep = report["users"] = {"created": 0, "existing": 0}
    for name in sorted(n for n in names if n):
        if name in users:
            rep["existing"] += 1
            continue
        u = User(name=name, role=role)  # 没有微信身份，之后由管理员并入其微信账号
        db.add(u)
        users[name] = u
        rep["created"] += 1
    db.flush()
    return users


def import_directory(db: Session, rows: list[dict], report: dict) -> dict[tuple[str, str], StoreDirectory]:
    rep = report["directory"] = {"total": len(rows), "created": 0, "updated": 0, "skipped": 0}
    existing = {(d.platform, d.external_id): d for d in db.scalars(select(StoreDirectory))}
    months = [c for c in (rows[0] if rows else {}) if re.fullmatch(r"\d{2}年\d{2}月", c)]
    for r in rows:
        platform, ext = s(r["销售平台"]), s(r["门店ID"])
        if not platform or not ext:
            rep["skipped"] += 1
            continue
        d = existing.get((platform, ext))
        if d is None:
            d = StoreDirectory(platform=platform, external_id=ext)
            db.add(d)
            existing[(platform, ext)] = d
            rep["created"] += 1
        else:
            rep["updated"] += 1
        d.name, d.province, d.city, d.district = s(r["门店名称"]), s(r.get("门店地址省")), s(r.get("门店地址市")), s(r.get("门店地址区"))
        d.address, d.grid, d.status = s(r.get("门店详细地址")), s(r.get("网格")), s(r.get("门店状态"))
        d.wecom_added, d.reported = (True if s(r.get("是否添加企微")) == "是" else None), (True if s(r.get("是否提报门店清单")) == "是" else None)
        d.el_code, d.service_station = s(r.get("门店EL码")), s(r.get("挂靠服务站") or r.get("挂靠服务站1"))
        d.gift_tier = s(r.get("赠品坎级"))
        d.monthly_sales = {f"20{m[:2]}-{m[3:5]}": int(num(r.get(m)) or 0) for m in months if num(r.get(m)) is not None}
    db.flush()
    return existing


def import_stores(db: Session, rows: list[dict], users: dict[str, User], directory: dict, report: dict) -> dict[str, Store]:
    rep = report["stores"] = {
        "total": len(rows), "created": 0, "updated": 0, "external_id_placeholder": 0, "external_id_conflicts": [],
        "linked_directory": 0, "no_directory_match": 0, "unknown_cooperation_status": Counter(), "profile_sections": 0,
    }
    all_stores = list(db.scalars(select(Store)))
    by_code = {st.code: st for st in all_stores if st.code}
    used_keys = {(st.platform, st.external_id): st for st in all_stores if st.platform and st.external_id}
    for r in rows:
        code = s(r["门店编号"])
        if not code:
            continue
        st = by_code.get(code)
        if st is None:
            st = Store(code=code, name=s(r["门店名称"]) or code)
            db.add(st)
            by_code[code] = st
            rep["created"] += 1
        else:
            rep["updated"] += 1
        st.name = s(r["门店名称"]) or st.name
        platform = s(r["装机平台"])
        st.platform = None if platform in ("", "无平台") else platform
        ext = s(r["门店ID"])
        if ext in ("无", "-", "/", "无ID"):
            ext = ""
            rep["external_id_placeholder"] += 1
        st.external_id = ext or None
        if st.platform and st.external_id:  # 同一平台的同一个 ID 只能属于一家店，后出现的清空并记下来
            owner = used_keys.get((st.platform, st.external_id))
            if owner is not None and owner is not st:
                rep["external_id_conflicts"].append(f"{code} {st.name}（与 {owner.code} {owner.name} 共用 {st.platform}/{st.external_id}）")
                st.external_id = None
            else:
                used_keys[(st.platform, st.external_id)] = st
        d = directory.get((st.platform, st.external_id)) if st.platform and st.external_id else None
        st.directory_id = d.id if d else None
        if st.platform and st.external_id:
            rep["linked_directory" if d else "no_directory_match"] += 1
        status = s(r["合作状态"]) or "未触达"
        if status not in COOPERATION_STATUSES:
            rep["unknown_cooperation_status"][status] += 1
        st.cooperation_status = status
        st.store_type = s(r.get("门店类型"))
        st.primary_sales_id = users[s(r["归属销售"])].id if s(r["归属销售"]) in users else None
        location = s(r.get("地理位置"))
        st.address = s(r.get("详细地址")) or (location.split(",", 1)[1].strip() if "," in location else location)
        st.province, st.grid = s(r.get("门店所在省")), s(r.get("网格"))
        st.contact_name, st.contact_phone = s(r.get("联系人姓名")), s(r.get("联系电话"))
        st.installed_at = parse_date(r.get("装机时间"))
        st.wecom_added = True if s(r.get("是否添加企微")) == "是" else None
        st.photo_key = s(r.get("门头照"))[:255]
    db.flush()

    # 档案七维度：内容来自拆开的各列，状态（稳定档案/当前状态/未确认）来自「门店档案」总文字
    existing = {(x.store_id, x.key): x for x in db.scalars(select(StoreProfileSection))}
    for r in rows:
        st = by_code.get(s(r["门店编号"]))
        if st is None:
            continue
        whole = s(r.get("门店档案"))
        for col, key in PROFILE_COLUMNS.items():
            content = s(r.get(col))
            if not content:
                continue
            m = re.search(rf"^{re.escape(PROFILE_LABEL_BY_KEY[key])}（(稳定档案|当前状态|未确认)）：", whole, re.M)
            sec = existing.get((st.id, key))
            if sec is None:
                sec = StoreProfileSection(store_id=st.id, key=key)
                db.add(sec)
                existing[(st.id, key)] = sec
            sec.state, sec.content = (m.group(1) if m else "当前状态"), content
            rep["profile_sections"] += 1
    db.flush()
    return by_code


def import_visits(db: Session, rows: list[dict], stores: dict[str, Store], users: dict[str, User], tz: ZoneInfo | None,
                  report: dict) -> dict[str, Visit]:
    price = get_settings().asr_price_per_hour
    rep = report["visits"] = {
        "total": len(rows), "created": 0, "updated": 0, "unknown_store": [], "stage_defaulted": 0, "stage_mapped": Counter(),
        "status": Counter(), "unknown_store_condition": Counter(), "segments": 0, "with_recording_but_no_transcript": 0,
        "transcripts_parsed": 0, "transcripts_raw_only": 0, "analysis_rows": 0, "dsr_corrections": 0, "no_time": 0,
    }
    hours = report["entered_hours_beijing"] = Counter()
    by_code = {v.code: v for v in db.scalars(select(Visit)) if v.code}
    seg_visit_ids = set(db.scalars(select(VisitSegment.visit_id).distinct()))
    corr_keys = {(c.store_id, c.visit_id) for c in db.scalars(select(StoreCorrection).where(StoreCorrection.source == "import"))}
    bj = ZoneInfo("Asia/Shanghai")

    for r in rows:
        code = s(r["拜访编号"])
        store = stores.get(s(r["门店编号"]))
        if not code or store is None:
            rep["unknown_store"].append(f"{code or '（无编号）'} → 门店编号 {s(r['门店编号'])}")
            continue
        v = by_code.get(code)
        if v is None:
            v = Visit(code=code, store_id=store.id, stage=FIRST_STAGE)
            db.add(v)
            by_code[code] = v
            rep["created"] += 1
        else:
            rep["updated"] += 1
        v.store_id, v.legacy = store.id, True
        visitor = users.get(s(r["拜访区域经理"]))
        v.visitor_id = visitor.id if visitor else None
        v.manager_id = None  # 旧数据只有一个人，不知道当时的上级
        entered, left = to_utc(r.get("进店时间"), tz), to_utc(r.get("离店时间"), tz)
        if entered is None:
            rep["no_time"] += 1
            entered = datetime(2000, 1, 1)
        v.entered_at, v.left_at = entered, left
        hours[entered.replace(tzinfo=timezone.utc).astimezone(bj).hour] += 1

        purposes = [p.strip() for p in re.split(r"[,，、]", s(r.get("本次拜访目的"))) if p.strip()]
        v.purposes = purposes or None
        raw_stage = s(r["拜访阶段1"])
        stage = LEGACY_STAGE_MAP.get(raw_stage, raw_stage)
        if stage not in STAGES:
            stage = FIRST_STAGE if "破冰建联" in purposes else "日常维护"
            rep["stage_defaulted"] += 1
        elif raw_stage in LEGACY_STAGE_MAP:
            rep["stage_mapped"][f"{raw_stage}→{stage}"] += 1
        v.stage = stage
        v.cooperated = s(r.get("是否达成合作")) or None
        v.note = s(r.get("现场速记"))
        cond = s(r.get("门店匹配状态"))
        v.store_condition = cond or None
        if cond and cond not in STORE_CONDITIONS:
            rep["unknown_store_condition"][cond] += 1
        v.checkin_address, v.photo_key = s(r.get("定位打卡"))[:255], s(r.get("门头照"))[:255]
        survey = {key: (num(r.get(col)) if key.endswith(("Total", "Est")) else s(r.get(col))) for col, key in SURVEY_COLUMNS.items()
                  if s(r.get(col))}
        v.survey = survey or None

        files = [f.strip() for f in s(r.get("录音文件")).split(",") if f.strip()]
        mode_none = s(r.get("录音方式")) == "无录音"
        v.recording_mode = "none" if mode_none else "uploaded"
        v.no_recording_reason = s(r.get("无录音原因"))
        minutes = num(r.get("音频时长（分钟）"))
        v.duration_sec = round(minutes * 60) if minutes else 0
        v.est_cost = round(v.duration_sec / 3600 * price, 2)

        # 转写
        text = s(r.get("完整对话-文字"))
        db.flush()
        if text:
            utterances = parse_transcript(text)
            upsert_transcript(db, v.id, utterances, raw_text=text, source="import")
            rep["transcripts_parsed" if utterances else "transcripts_raw_only"] += 1

        # 分析结果（原文）
        any_analysis = False
        for col, module, evidence_col in ANALYSIS_COLUMNS:
            body = s(r.get(col))
            if not body:
                continue
            upsert_module(db, v.id, module, {"text": body}, evidence=s(r.get(evidence_col)) if evidence_col else None)
            rep["analysis_rows"] += 1
            any_analysis = True

        # 状态
        validity = s(r.get("是否有效对话"))
        if mode_none:
            v.status = "no_recording"
        elif validity == "内容无效":
            v.status = "invalid_content"
        elif validity == "录音过短":
            v.status = "invalid_short"
        elif any_analysis:
            v.status = "done"
        elif files:
            v.status = "cost_pending"  # 有录音但还没转写，等 AI 接入后处理
            rep["with_recording_but_no_transcript"] += 1
        else:
            v.status = "failed"
        rep["status"][v.status] += 1

        # 录音文件（只有文件名，音频另行搬运）
        if files and v.id not in seg_visit_ids:
            for i, name in enumerate(files):
                ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
                db.add(VisitSegment(visit_id=v.id, seq=i, object_key=f"legacy/{code}/{i:03d}-{name}"[:255], original_name=name[:255],
                                    format=ext[:8], file_missing=True,
                                    duration_ms=v.duration_sec * 1000 if len(files) == 1 else 0))
                rep["segments"] += 1

        # DSR 手填的档案纠正
        dsr = s(r.get("门店档案纠正（dsr填写）"))
        if dsr and (store.id, v.id) not in corr_keys:
            db.add(StoreCorrection(store_id=store.id, visit_id=v.id, user_id=v.visitor_id, text=dsr, source="import"))
            corr_keys.add((store.id, v.id))
            rep["dsr_corrections"] += 1
    db.flush()
    return by_code


def link_profile_sources(db: Session, store_rows: list[dict], stores: dict[str, Store], visits: dict[str, Visit]) -> None:
    """档案维度的「来源拜访」：取主档里关联的拜访中最近的一次"""
    sections = {}
    for sec in db.scalars(select(StoreProfileSection)):
        sections.setdefault(sec.store_id, []).append(sec)
    for r in store_rows:
        st = stores.get(s(r["门店编号"]))
        linked = [visits[c] for c in re.findall(r"RA-\d+", s(r.get("拜访记录"))) if c in visits]
        if st is None or not linked:
            continue
        latest = max(linked, key=lambda v: v.entered_at)
        for sec in sections.get(st.id, []):
            sec.source_visit_id = latest.id


def import_todos(db: Session, rows: list[dict], stores: dict[str, Store], year: int, report: dict) -> None:
    """旧「后续行动」表其实是月度目标跟踪：X月目标 / 目前达成 / 跟进情况"""
    rep = report["todos"] = {"total": len(rows), "created": 0, "skipped_existing": 0, "store_unmatched": [], "store_ambiguous": []}
    by_name: dict[str, list[Store]] = {}
    for st in stores.values():
        by_name.setdefault(st.name, []).append(st)
    existing = {(t.store_id, t.period, t.topic) for t in db.scalars(select(Todo).where(Todo.source == "import"))}
    for r in rows:
        month_col = next((c for c in r if re.fullmatch(r"\d{1,2}月目标", c)), None)
        month = int(month_col.split("月")[0]) if month_col else None
        period = f"{year}-{month:02d}" if month else None
        name = s(r.get("门店名称"))
        matches = by_name.get(name, [])
        store = matches[0] if len(matches) == 1 else None
        if name and not matches:
            rep["store_unmatched"].append(name)
        elif len(matches) > 1:
            rep["store_ambiguous"].append(name)
        topic = s(r["行动事项"]) or "任务"
        if store is None and name:
            topic = f"{name}：{topic}"
        key = (store.id if store else None, period, topic)
        if key in existing:
            rep["skipped_existing"] += 1
            continue
        due = None
        if month:
            nxt = date(year + (month == 12), month % 12 + 1, 1)
            due = nxt - timedelta(days=1)
        db.add(Todo(
            source="import", store_id=store.id if store else None, assignee_id=store.primary_sales_id if store else None,
            topic=topic[:128], period=period, due_date=due, target_qty=num(r.get(month_col)) if month_col else None,
            achieved_qty=num(r.get("目前达成")), progress_note=s(r.get("跟进情况"))))
        existing.add(key)
        rep["created"] += 1
    db.flush()


# ---------- 总入口 ----------
def run(db: Session, wb, parts: list[str], tz_name: str, role: str, year: int | None) -> dict:
    tz = None if tz_name.upper() in ("NONE", "") else ZoneInfo(tz_name)
    report: dict = {"warnings": []}
    dir_rows = read_sheet(wb, SHEET_DIRECTORY, report) if "directory" in parts else []
    store_rows = read_sheet(wb, SHEET_STORES, report) if "stores" in parts or "visits" in parts else []
    visit_rows = read_sheet(wb, SHEET_VISITS, report) if "visits" in parts else []
    todo_rows = read_sheet(wb, SHEET_TODOS, report) if "todos" in parts else []

    names = {s(r.get("归属销售")) for r in store_rows} | {s(r.get("拜访区域经理")) for r in visit_rows}
    users = get_users(db, names, role, report)
    directory = import_directory(db, dir_rows, report) if dir_rows else {(d.platform, d.external_id): d for d in db.scalars(select(StoreDirectory))}
    stores = import_stores(db, store_rows, users, directory, report) if "stores" in parts else {st.code: st for st in db.scalars(select(Store)) if st.code}
    visits: dict[str, Visit] = {}
    if "visits" in parts:
        visits = import_visits(db, visit_rows, stores, users, tz, report)
        link_profile_sources(db, store_rows, stores, visits)
    if "todos" in parts and todo_rows:
        y = year or (max((v.entered_at.year for v in visits.values()), default=datetime.now().year))
        import_todos(db, todo_rows, stores, y, report)
    return report


def print_report(report: dict, committed: bool) -> None:
    def show(title: str, data) -> None:
        print(f"\n【{title}】")
        for k, v in data.items():
            if isinstance(v, (Counter, dict)):
                v = dict(v)
            if isinstance(v, list):
                print(f"  {k}: {len(v)} 条" + (f"，前 5 条：{v[:5]}" if v else ""))
            else:
                print(f"  {k}: {v}")

    for key, title in (("users", "人员"), ("directory", "总门店清单"), ("stores", "门店主档"), ("visits", "拜访"), ("todos", "后续行动（月度目标）")):
        if key in report:
            show(title, report[key])
    if report.get("entered_hours_beijing"):
        h = sorted(report["entered_hours_beijing"].items())
        working = sum(n for hour, n in h if 6 <= hour <= 20)
        total = sum(n for _, n in h)
        print(f"\n【时区自查】按北京时间看进店时间，落在 6～20 点的占 {working}/{total}。"
              f"如果比例很低，说明 Excel 里的时间已经是北京时间，请加参数 --tz Asia/Shanghai 重新试运行。")
    for w in report["warnings"]:
        print(f"\n警告：{w}")
    print("\n" + ("已写入数据库。" if committed else "这是试运行，没有写入任何数据。确认无误后加 --commit 真正导入。"))


def main() -> None:
    parser = argparse.ArgumentParser(description="把飞书导出的 Excel 导入数据库")
    parser.add_argument("--file", required=True, help="飞书多维表导出的 .xlsx")
    parser.add_argument("--commit", action="store_true", help="真正写库（默认只试运行）")
    parser.add_argument("--only", default=",".join(ALL_PARTS), help=f"只导入其中几部分，逗号分隔：{','.join(ALL_PARTS)}")
    parser.add_argument("--tz", default="UTC", help="Excel 里时间的时区，默认 UTC；如果已是北京时间填 Asia/Shanghai")
    parser.add_argument("--default-role", default="manager", choices=["sales", "manager"], help="新建人员的角色（旧表里的拜访人是区域经理）")
    parser.add_argument("--year", type=int, help="月度目标所属年份，默认取最近一次拜访的年份")
    parser.add_argument("--report", help="把报告另存为 JSON 文件")
    args = parser.parse_args()

    parts = [p.strip() for p in args.only.split(",") if p.strip()]
    bad = [p for p in parts if p not in ALL_PARTS]
    if bad:
        sys.exit(f"--only 里有不认识的部分：{bad}")
    import openpyxl  # 只有导入时才需要

    from app.db import SessionLocal

    wb = openpyxl.load_workbook(args.file, data_only=True)
    with SessionLocal() as db:
        try:
            report = run(db, wb, parts, args.tz, args.default_role, args.year)
            if args.commit:
                db.commit()
            else:
                db.rollback()
        except BaseException:
            db.rollback()
            raise
    print_report(report, args.commit)
    if args.report:
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=lambda o: dict(o)), encoding="utf-8")


if __name__ == "__main__":
    main()
