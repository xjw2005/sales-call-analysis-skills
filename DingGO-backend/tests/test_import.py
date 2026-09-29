"""导入脚本测试：用手工造的迷你 Excel（表头与飞书导出一致，内容是虚构的），不含任何真实数据"""

import importlib.util
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from app.db import SessionLocal
from app.models import Store, StoreCorrection, StoreDirectory, StoreProfileSection, Todo, User, Visit, VisitAnalysis, VisitSegment, VisitTranscript

from .conftest import ADMIN, login_user

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "import_feishu_xlsx.py"
spec = importlib.util.spec_from_file_location("import_feishu_xlsx", SCRIPT)
imp = importlib.util.module_from_spec(spec)
sys.modules["import_feishu_xlsx"] = imp
spec.loader.exec_module(imp)

PROFILE_TEXT = (
    "一句话画像：新开母婴店，老板无经验。\n\n门店基本信息（稳定档案）：夫妻店。\n\n主营品类与品牌（当前状态）：奶粉。\n\n"
    "经营模式（未确认）：未确认\n\n选品偏好（当前状态）：好卖的流通品。\n\n利润偏好（当前状态）：怕被打穿。\n\n"
    "合作偏好与排斥项（当前状态）：要红线价。\n\n档案口径：动态更新档案"
)
TRANSCRIPT = (
    "[00:11.450–00:12.290] 销售：老板您好。\n[00:34.490–00:37.940] 客户：你说吧。\n续行没有时间戳\n"
    "[01:00:03.410–01:00:07.430] 旁人：老板，这个放哪？"
)


def sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])


def build_workbook(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    months = [f"26年{m:02d}月" for m in range(1, 10)]
    sheet(wb, imp.SHEET_DIRECTORY,
          ["创建时间", "门店名称", "销售平台", "门店地址省", "门店地址市", "门店地址区", "门店ID", "是否添加企微", "门店EL码", "挂靠服务站",
           "门店状态", "网格", "赠品坎级", "门店详细地址", "联系人", "联系电话"] + months,
          [{"门店名称": "宝贝乐", "销售平台": "西港", "门店地址省": "重庆市", "门店地址市": "渝北区", "门店地址区": "龙溪", "门店ID": "9001",
            "是否添加企微": "是", "门店状态": "启用", "网格": "A1", "赠品坎级": "小于4罐", "门店详细地址": "新南路", **{m: i for i, m in enumerate(months)}},
           {"门店名称": "智生活店甲", "销售平台": "智生活", "门店ID": "7378", "门店状态": "启用", **{m: 0 for m in months}},
           {"门店名称": "智生活店乙", "销售平台": "智生活", "门店ID": "7410", "门店状态": "禁用", **{m: 0 for m in months}}])
    store_header = ["门店编号", "门店名称", "装机平台", "门店ID", "门店类型", "合作状态", "归属销售", "详细地址", "门头照", "地理位置", "门店所在省",
                    "网格", "联系人姓名", "联系电话", "装机时间", "是否添加企微", "拜访记录", "门店档案", "一句画像", "门店基本信息",
                    "主营品类与品牌", "经营模式", "选品偏好", "利润偏好", "合作偏好与排斥项"]
    sheet(wb, imp.SHEET_STORES, store_header, [
        {"门店编号": "ST-0001", "门店名称": "宝贝乐母婴店", "装机平台": "西港", "门店ID": "9001", "门店类型": "系统门店", "合作状态": "已合作",
         "归属销售": "赵经理", "详细地址": "新南路 88 号", "门头照": "IMG_1.jpg", "门店所在省": "重庆市", "网格": "A1", "联系人姓名": "王老板",
         "联系电话": "13800000000", "装机时间": "2025-03-05", "是否添加企微": "是", "拜访记录": "RA-0001, RA-0003", "门店档案": PROFILE_TEXT,
         "一句画像": "新开母婴店，老板无经验。", "门店基本信息": "夫妻店。", "主营品类与品牌": "奶粉。", "经营模式": "未确认",
         "选品偏好": "好卖的流通品。", "利润偏好": "怕被打穿。", "合作偏好与排斥项": "要红线价。"},
        {"门店编号": "ST-0002", "门店名称": "无平台小店", "装机平台": "无平台", "门店ID": "无", "合作状态": "已触达未合作", "归属销售": "钱经理",
         "地理位置": "无平台小店, 云南省昆明市官渡区某街 1 号"},
        {"门店编号": "ST-0003", "门店名称": "智生活店甲", "装机平台": "智生活", "门店ID": "7378", "合作状态": "已合作", "归属销售": "赵经理"},
        {"门店编号": "ST-0004", "门店名称": "智生活店甲（重复登记）", "装机平台": "智生活", "门店ID": "7378", "合作状态": None, "归属销售": "赵经理"},
        {"门店编号": "ST-0005", "门店名称": "重名店", "装机平台": "西港", "门店ID": None, "合作状态": "未触达", "归属销售": "赵经理"},
        {"门店编号": "ST-0006", "门店名称": "重名店", "装机平台": "西港", "门店ID": None, "合作状态": "未触达", "归属销售": "赵经理"},
    ])
    visit_header = ["拜访编号", "门店编号", "拜访区域经理", "进店时间", "离店时间", "录音方式", "音频时长（分钟）", "门店匹配状态", "门头照", "无录音原因",
                    "定位打卡", "录音文件", "完整对话-文字", "拜访阶段1", "本次拜访目的", "现场速记", "是否达成合作", "是否有效对话", "门店档案纠正（dsr填写）",
                    "AI拜访摘要", "显性需求(仅供参考)", "显性需求_原句参考", "门店档案", "门店档案-原文证据", "下一步行动策略",
                    "上一次行动与这一次行动总结闭环", "店铺面积", "店铺全品类月均销售（罐）", "是否做跨境", "a2月均预估销售（罐）"]
    sheet(wb, imp.SHEET_VISITS, visit_header, [
        {"拜访编号": "RA-0001", "门店编号": "ST-0001", "拜访区域经理": "赵经理", "进店时间": datetime(2026, 8, 3, 1, 30), "离店时间": datetime(2026, 8, 3, 2, 10),
         "录音方式": "已上传音频文件", "音频时长（分钟）": "10.5", "门店匹配状态": "正常运营", "定位打卡": "重庆市，渝北区·新南路", "录音文件": "1785738863.ogg,货品讨论.m4a",
         "完整对话-文字": TRANSCRIPT, "拜访阶段1": "陌拜破冰", "本次拜访目的": "破冰建联, 推进装机", "现场速记": "老板在店", "是否达成合作": "否", "是否有效对话": "有效",
         "门店档案纠正（dsr填写）": "门店约 80 平", "AI拜访摘要": "摘要文字", "显性需求(仅供参考)": "1. 需求点：陈列支持", "显性需求_原句参考": "[00:34] 客户：你说吧",
         "门店档案": PROFILE_TEXT, "门店档案-原文证据": "证据文字", "下一步行动策略": "周五前送物料", "店铺面积": "30-50平", "店铺全品类月均销售（罐）": 120,
         "是否做跨境": "是", "a2月均预估销售（罐）": 30},
        {"拜访编号": "RA-0002", "门店编号": "ST-0002", "拜访区域经理": "钱经理", "进店时间": datetime(2026, 8, 4, 3, 0), "离店时间": datetime(2026, 8, 4, 3, 20),
         "录音方式": "无录音", "无录音原因": "对方不同意", "门店匹配状态": "已闭店", "拜访阶段1": "日常拜访", "本次拜访目的": "催动销", "是否达成合作": "否"},
        {"拜访编号": "RA-0003", "门店编号": "ST-0001", "拜访区域经理": "赵经理", "进店时间": datetime(2026, 8, 20, 4, 0), "离店时间": datetime(2026, 8, 20, 4, 5),
         "录音方式": "已上传音频文件", "音频时长（分钟）": "1.2", "录音文件": "short.m4a", "拜访阶段1": "日常拜访", "是否有效对话": "录音过短", "门店匹配状态": "奇怪状况"},
        {"拜访编号": "RA-0004", "门店编号": "ST-9999", "拜访区域经理": "赵经理", "进店时间": datetime(2026, 8, 21, 4, 0), "拜访阶段1": "日常拜访"},
        {"拜访编号": "RA-0005", "门店编号": "ST-0003", "拜访区域经理": "赵经理", "进店时间": datetime(2026, 8, 22, 4, 0), "离店时间": datetime(2026, 8, 22, 4, 30),
         "录音方式": "已上传音频文件", "音频时长（分钟）": "20", "录音文件": "raw.ogg", "拜访阶段1": None, "本次拜访目的": "破冰建联"},
    ])
    sheet(wb, imp.SHEET_TODOS, ["行动事项", "任务来源", "完成时间", "门店名称", "9月目标", "目前达成", "差额", "跟进情况"], [
        {"行动事项": "9月任务拆解", "门店名称": "宝贝乐母婴店", "9月目标": "60", "目前达成": "36", "差额": "24", "跟进情况": "月中下单"},
        {"行动事项": "9月任务拆解", "门店名称": "查无此店", "9月目标": "24", "目前达成": "18", "差额": "6", "跟进情况": "待跟进"},
        {"行动事项": "9月任务拆解", "门店名称": "重名店", "9月目标": "10", "目前达成": "0", "差额": "10"},
    ])
    wb.save(path)


@pytest.fixture
def xlsx(tmp_path):
    p = tmp_path / "feishu.xlsx"
    build_workbook(p)
    return p


def do_import(xlsx, commit=True, **kw):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    with SessionLocal() as db:
        report = imp.run(db, wb, kw.get("parts", imp.ALL_PARTS), kw.get("tz", "UTC"), "manager", kw.get("year"))
        db.commit() if commit else db.rollback()
    return report


def count(model):
    with SessionLocal() as db:
        return db.query(model).count()


def test_parse_transcript():
    u = imp.parse_transcript(TRANSCRIPT)
    assert [x["role"] for x in u] == ["销售", "客户", "旁人"]
    assert u[0]["startMs"] == 11450 and u[0]["endMs"] == 12290 and u[2]["startMs"] == 3603410  # 超过 1 小时的时间戳
    assert u[1]["text"] == "你说吧。\n续行没有时间戳"  # 没有时间戳的行接在上一句后面
    assert u[0]["uid"] == "U0001"


def test_dry_run_writes_nothing(xlsx):
    report = do_import(xlsx, commit=False)
    assert report["stores"]["created"] == 6 and report["visits"]["created"] == 4
    assert all(count(m) == 0 for m in (User, Store, Visit, StoreDirectory, Todo))


def test_full_import(xlsx):
    report = do_import(xlsx)
    assert report["users"] == {"created": 2, "existing": 0}
    assert report["directory"]["created"] == 3
    st = report["stores"]
    assert st["created"] == 6 and st["external_id_placeholder"] == 1 and len(st["external_id_conflicts"]) == 1
    assert "ST-0004" in st["external_id_conflicts"][0] and st["linked_directory"] == 2
    assert dict(st["unknown_cooperation_status"]) == {} and st["profile_sections"] == 7

    with SessionLocal() as db:
        users = {u.name: u for u in db.query(User)}
        assert users["赵经理"].role == "manager" and users["赵经理"].openid is None
        s1 = db.query(Store).filter_by(code="ST-0001").one()
        assert (s1.platform, s1.external_id, s1.cooperation_status, s1.contact_phone) == ("西港", "9001", "已合作", "13800000000")
        assert s1.directory_id is not None and s1.primary_sales_id == users["赵经理"].id and s1.installed_at.isoformat() == "2025-03-05"
        assert s1.wecom_added is True and s1.photo_key == "IMG_1.jpg"
        s2 = db.query(Store).filter_by(code="ST-0002").one()
        assert s2.platform is None and s2.external_id is None and s2.address == "云南省昆明市官渡区某街 1 号"  # 无平台、ID 为「无」；地址取自地理位置
        s4 = db.query(Store).filter_by(code="ST-0004").one()
        assert s4.external_id is None and s4.cooperation_status == "未触达"  # 重复的 ID 被清空
        # 档案：内容来自拆开的列，状态来自总文字
        secs = {x.key: x for x in db.query(StoreProfileSection).filter_by(store_id=s1.id)}
        assert secs["basic"].state == "稳定档案" and secs["business_model"].state == "未确认" and secs["one_line"].content.startswith("新开母婴店")
        assert secs["basic"].source_visit_id == db.query(Visit).filter_by(code="RA-0003").one().id  # 来源 = 关联拜访里最近的

        v1 = db.query(Visit).filter_by(code="RA-0001").one()
        assert v1.legacy and v1.stage == "首访破冰" and v1.status == "done" and v1.recording_mode == "uploaded"
        assert v1.entered_at == datetime(2026, 8, 3, 1, 30) and v1.duration_sec == 630 and v1.cooperated == "否"
        assert v1.purposes == ["破冰建联", "推进装机"] and v1.store_condition == "正常运营" and v1.note == "老板在店"
        assert v1.survey == {"area": "30-50平", "monthlyTotal": 120.0, "crossBorder": "是", "a2MonthlyEst": 30.0}
        assert v1.checkin_address == "重庆市，渝北区·新南路" and v1.visitor_id == users["赵经理"].id and v1.manager_id is None
        segs = db.query(VisitSegment).filter_by(visit_id=v1.id).order_by(VisitSegment.seq).all()
        assert [(x.original_name, x.format, x.file_missing) for x in segs] == [("1785738863.ogg", "ogg", True), ("货品讨论.m4a", "m4a", True)]
        tr = db.get(VisitTranscript, v1.id)
        assert len(tr.utterances) == 3 and tr.source == "import" and tr.raw_text.startswith("[00:11")
        mods = {a.module: a for a in db.query(VisitAnalysis).filter_by(visit_id=v1.id)}
        assert set(mods) == {"ai-summary", "explicit-needs", "store-profile", "next-action"}
        assert mods["explicit-needs"].result == {"text": "1. 需求点：陈列支持"} and mods["explicit-needs"].evidence.startswith("[00:34]")
        assert db.query(StoreCorrection).filter_by(visit_id=v1.id).one().text == "门店约 80 平"

        v2 = db.query(Visit).filter_by(code="RA-0002").one()
        assert (v2.status, v2.recording_mode, v2.no_recording_reason, v2.stage, v2.store_condition) == ("no_recording", "none", "对方不同意", "日常维护", "已闭店")
        v3 = db.query(Visit).filter_by(code="RA-0003").one()
        assert v3.status == "invalid_short" and v3.store_condition == "奇怪状况"
        v5 = db.query(Visit).filter_by(code="RA-0005").one()
        assert v5.status == "cost_pending" and v5.stage == "首访破冰" and v5.duration_sec == 1200 and v5.est_cost > 0  # 阶段为空：目的含「破冰建联」→ 首访
        assert db.query(Visit).filter_by(code="RA-0004").first() is None

    vr = report["visits"]
    assert vr["created"] == 4 and len(vr["unknown_store"]) == 1 and vr["stage_defaulted"] == 1
    assert dict(vr["stage_mapped"]) == {"陌拜破冰→首访破冰": 1, "日常拜访→日常维护": 2} and dict(vr["unknown_store_condition"]) == {"奇怪状况": 1}
    assert dict(vr["status"]) == {"done": 1, "no_recording": 1, "invalid_short": 1, "cost_pending": 1} and vr["with_recording_but_no_transcript"] == 1
    assert vr["segments"] == 4 and vr["dsr_corrections"] == 1 and vr["transcripts_parsed"] == 1
    assert report["entered_hours_beijing"][9] == 1  # 01:30 UTC = 北京 9 点

    tr = report["todos"]
    assert tr["created"] == 3 and tr["store_unmatched"] == ["查无此店"] and tr["store_ambiguous"] == ["重名店"]
    with SessionLocal() as db:
        t = db.query(Todo).filter_by(store_id=db.query(Store).filter_by(code="ST-0001").one().id).one()
        assert (t.source, t.period, t.target_qty, t.achieved_qty, t.progress_note, t.due_date.isoformat()) == ("import", "2026-09", 60.0, 36.0, "月中下单", "2026-09-30")
        assert t.assignee_id is not None
        assert db.query(Todo).filter(Todo.store_id.is_(None)).count() == 2 and db.query(Todo).filter(Todo.topic.like("查无此店%")).count() == 1


def test_import_is_idempotent(xlsx):
    do_import(xlsx)
    before = {m.__name__: count(m) for m in (User, Store, Visit, StoreDirectory, StoreProfileSection, VisitSegment, VisitTranscript, VisitAnalysis, StoreCorrection, Todo)}
    report = do_import(xlsx)
    after = {m.__name__: count(m) for m in (User, Store, Visit, StoreDirectory, StoreProfileSection, VisitSegment, VisitTranscript, VisitAnalysis, StoreCorrection, Todo)}
    assert before == after
    assert report["stores"]["created"] == 0 and report["stores"]["updated"] == 6 and report["visits"]["created"] == 0
    assert report["todos"]["created"] == 0 and report["todos"]["skipped_existing"] == 3


def test_timezone_option(xlsx):
    report = do_import(xlsx, tz="Asia/Shanghai")
    with SessionLocal() as db:
        assert db.query(Visit).filter_by(code="RA-0001").one().entered_at == datetime(2026, 8, 2, 17, 30)  # 北京 1:30 = UTC 前一天 17:30
    assert report["entered_hours_beijing"][1] == 1


def test_missing_required_column(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb, imp.SHEET_STORES, ["门店编号", "门店名称"], [{"门店编号": "ST-1", "门店名称": "x"}])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(SystemExit) as e:
        do_import(p, parts=["stores"])
    assert "缺少必需的列" in str(e.value)


def test_imported_data_visible_through_api(client, xlsx):
    """导入后，旧人员并入微信账号，就能通过接口看到自己的门店、拜访（原文分析）和待办"""
    do_import(xlsx)
    with SessionLocal() as db:
        zhao_id = db.query(User).filter_by(name="赵经理").one().id
    h, new = login_user(client, "zhao-wx", "赵")
    assert client.post("/admin/users/merge", json={"fromId": new["id"], "intoId": str(zhao_id)}, headers=ADMIN).status_code == 200
    h, me = login_user(client, "zhao-wx")
    assert me["id"] == str(zhao_id) and me["role"] == "manager"

    stores = client.get("/stores", headers=h).json()
    by_code = {s["code"]: s for s in stores}
    assert set(by_code) == {"ST-0001", "ST-0003", "ST-0004", "ST-0005", "ST-0006"}  # 钱经理的门店不可见
    assert by_code["ST-0001"]["cooperationStatus"] == "已合作" and "contactPhone" not in by_code["ST-0001"]
    assert by_code["ST-0001"]["oneLine"].startswith("新开母婴店") and by_code["ST-0001"]["visitCount"] == 2
    detail = client.get(f"/stores/{by_code['ST-0001']['id']}", headers=h).json()
    assert detail["contactPhone"] == "13800000000" and detail["correction"] == "门店约 80 平"
    assert [s["title"] for a in detail["legacyActions"] for s in a["sections"]] == ["下一步行动"] and len(detail["legacyActions"]) == detail["visitCount"] and isinstance(detail["todos"], list)

    visits = client.get("/visits", headers=h).json()
    assert {v["code"] for v in visits} == {"RA-0001", "RA-0003", "RA-0005"} and all(v["legacy"] for v in visits)
    v1 = next(v for v in visits if v["code"] == "RA-0001")
    full = client.get(f"/visits/{v1['id']}", headers=h).json()
    assert [x["title"] for x in full["analysis"]["legacySections"]] == ["AI 拜访摘要", "显性需求", "门店档案", "下一步行动策略"]
    assert len(full["transcript"]) == 3 and full["audioUrls"] == [] and full["audioMissing"] == 2  # 音频还没搬过来

    # 历史拜访的录音文件还没迁移：不能确认识别
    v5 = next(v for v in visits if v["code"] == "RA-0005")
    refuse = client.post(f"/visits/{v5['id']}/confirm-cost", headers=h)
    assert refuse.status_code == 409 and "还没有迁移" in refuse.json()["message"]

    brief = client.get(f"/assistant/brief/{by_code['ST-0001']['id']}", headers=h).json()
    assert brief["isFirst"] is False and any("补全「经营模式」" in q for q in brief["questions"])  # 用导入的档案生成要问的问题
    todos = client.get("/todos", headers=h).json()
    assert [t["gapQty"] for t in todos if t["storeName"] == "宝贝乐母婴店"] == [24.0]
    # 今日待办里，目标类待办显示进度
    panel = client.get(f"/assistant/today?storeId={by_code['ST-0001']['id']}", headers=h).json()
    rows = [r for sec in panel["sections"] if sec["title"] == "跟进事项" for r in sec["rows"]]
    assert any("目标 60，已达成 36，差 24" in r["sub"] for r in rows)
